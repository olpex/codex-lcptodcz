from pathlib import Path

import pytest
from openpyxl import Workbook

from app.models import Document, DocumentType, ExportJob, ImportJob, JobStatus, Trainee
from app.tasks.worker import process_export_job_task, process_import_job_task


def _create_import_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["first_name", "last_name", "status"])
    sheet.append(["Марина", "Іваненко", "active"])
    workbook.save(path)


def test_import_job_lifecycle_success(db_session, tmp_path: Path):
    file_path = tmp_path / "import.xlsx"
    _create_import_xlsx(file_path)

    document = Document(
        file_name="import.xlsx",
        file_path=str(file_path),
        file_type=DocumentType.XLSX,
        source="upload",
    )
    db_session.add(document)
    db_session.flush()

    job = ImportJob(
        idempotency_key="it-import-1",
        document_id=document.id,
        status=JobStatus.QUEUED,
        message="queued",
    )
    db_session.add(job)
    db_session.commit()

    result = process_import_job_task.run(job.id)
    assert result["status"] == "ok"

    db_session.refresh(job)
    assert job.status == JobStatus.SUCCEEDED
    assert job.result_payload is not None
    assert job.result_payload["import_result"]["inserted"] == 1

    trainees = db_session.query(Trainee).filter(Trainee.first_name == "Марина").all()
    assert len(trainees) == 1


def test_import_job_marks_failed_and_increments_retry(db_session):
    document = Document(
        file_name="missing.xlsx",
        file_path="C:/no/such/file.xlsx",
        file_type=DocumentType.XLSX,
        source="upload",
    )
    db_session.add(document)
    db_session.flush()

    job = ImportJob(
        idempotency_key="it-import-fail-1",
        document_id=document.id,
        status=JobStatus.QUEUED,
        message="queued",
    )
    db_session.add(job)
    db_session.commit()

    with pytest.raises(Exception):
        process_import_job_task.run(job.id)

    db_session.refresh(job)
    assert job.status == JobStatus.FAILED
    assert (job.retries or 0) >= 1
    assert job.message


def test_export_job_lifecycle_success(db_session):
    db_session.add(Trainee(first_name="Ірина", last_name="Козак", status="active"))
    db_session.commit()

    job = ExportJob(
        idempotency_key="it-export-1",
        report_type="trainees",
        export_format="csv",
        status=JobStatus.QUEUED,
        message="queued",
    )
    db_session.add(job)
    db_session.commit()

    result = process_export_job_task.run(job.id)
    assert result["status"] == "ok"

    db_session.refresh(job)
    assert job.status == JobStatus.SUCCEEDED
    assert job.output_document_id is not None
    assert job.result_payload is not None
    assert job.result_payload["rows"] >= 1
